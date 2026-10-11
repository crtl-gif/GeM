"""Optional Streamlit workbench for document extraction and AI-assisted tender review."""
import io, os
import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from ai_router import AIRouter
load_dotenv()
st.set_page_config(page_title="GeM Tender Assistant",page_icon="📄",layout="wide")
st.title("GeM Tender Assistant")
st.caption("Prepare and review tender documents. Human review and final submission are required.")
with st.sidebar:
 st.subheader("AI connection")
 configured=AIRouter().configured()
 st.write("Gemini:", "Configured" if configured["Gemini"] else "No key")
 st.write("Groq:", "Configured" if configured["Groq"] else "No key")
 st.warning("Tender text sent to a cloud AI provider leaves your computer. Upload only if authorised.")
def extract(file):
 name=file.name.lower(); data=file.getvalue()
 if name.endswith(".pdf"):
  import fitz
  doc=fitz.open(stream=data,filetype="pdf")
  return "\n".join(p.get_text() for p in doc)
 if name.endswith(".docx"):
  from docx import Document
  doc=Document(io.BytesIO(data))
  return "\n".join(p.text for p in doc.paragraphs)+"\n"+"\n".join(" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
 if name.endswith((".xlsx",".xls")):
  sheets=pd.read_excel(io.BytesIO(data),sheet_name=None)
  return "\n".join("SHEET: "+k+"\n"+v.to_string(index=False) for k,v in sheets.items())
 if name.endswith(".csv"): return pd.read_csv(io.BytesIO(data)).to_string(index=False)
 if name.endswith((".txt",".md")): return data.decode("utf-8",errors="replace")
 raise ValueError("Supported files: PDF, DOCX, XLSX, XLS, CSV, TXT, MD")
tab1,tab2=st.tabs(["Tender analysis","Document extraction"])
with tab1:
 uploaded=st.file_uploader("Upload tender / bid documents",type=["pdf","docx","xlsx","xls","csv","txt","md"],accept_multiple_files=True,key="analysis")
 manual=st.text_area("Or paste tender text",height=180)
 if st.button("Extract and analyse",type="primary"):
  parts=[]
  for f in uploaded or []:
   try: parts.append(f"DOCUMENT: {f.name}\n"+extract(f))
   except Exception as e: st.error(f"{f.name}: {e}")
  source="\n\n".join(parts+[manual])
  if not source.strip(): st.warning("Upload a document or paste text.")
  else:
   prompt="""You are a careful GeM tender preparation assistant. Analyse the supplied text and return: (1) short summary and important dates, (2) eligibility and mandatory conditions, (3) EMD amount and exemption conditions if stated, (4) technical/commercial requirements, (5) document checklist, (6) clarification questions, (7) risks and items requiring manual verification. Do not invent missing facts. Cite the clause or page where possible. Clearly flag uncertainty. This is preparation assistance, not legal advice.\n\nTEXT:\n"""+source[:100000]
   try:
    with st.spinner("Analysing with configured AI provider…"): result=AIRouter().generate(prompt)
    st.markdown(result); st.download_button("Download analysis",result,file_name="tender_analysis.md")
   except Exception as e: st.error(str(e))
with tab2:
 f=st.file_uploader("Choose one document",type=["pdf","docx","xlsx","xls","csv","txt","md"],key="extract")
 if f:
  try:
   text=extract(f); st.text_area("Extracted text",text,height=350)
   st.download_button("Download extracted text",text,file_name=f.name+".txt")
  except Exception as e: st.error(str(e))
st.divider()
st.caption("Verify every date, eligibility clause, EMD exemption and attachment against the official GeM bid. This app does not solve CAPTCHA or submit bids.")
